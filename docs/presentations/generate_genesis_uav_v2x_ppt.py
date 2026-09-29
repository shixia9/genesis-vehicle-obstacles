"""Generate the Genesis World UAV/V2X research-group presentation.

Run from the repository root:
    .venv/bin/python docs/presentations/generate_genesis_uav_v2x_ppt.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).with_name("genesis_uav_v2x_research_overview.pptx")
TEASER = ROOT / "imgs" / "teaser.png"
LOGO = ROOT / "imgs" / "logo_with_text.png"
ROBOT_IMAGE = ROOT / "out/mobile_robot_ground_truth_calibrated_v2/rgb_robot/frame_00025.png"

W, H = 13.333, 7.5

BG = "09111F"
PANEL = "111D2D"
PANEL_2 = "17263A"
WHITE = "F7FAFC"
MUTED = "A9B7CA"
CYAN = "28D7E6"
BLUE = "548BFF"
PURPLE = "9B7BFF"
GREEN = "4DE0A6"
YELLOW = "FFC857"
RED = "FF6B6B"
GRID = "26364C"

FONT = "PingFang SC"
MONO = "SFMono-Regular"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    *,
    size=18,
    color=WHITE,
    bold=False,
    font=FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.04,
    line_spacing=1.05,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.text = text
    p.alignment = align
    p.line_spacing = line_spacing
    # PowerPoint may split explicit line breaks into multiple runs. Style every
    # run so continuation lines do not fall back to the theme's black text.
    for run in p.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def add_rich_lines(slide, lines, x, y, w, h, *, size=17, color=WHITE, bullet=False, gap=5):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(0.06)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    for idx, item in enumerate(lines):
        if isinstance(item, tuple):
            text, item_color, item_bold = item
        else:
            text, item_color, item_bold = item, color, False
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = text
        p.level = 0
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = item_bold
        p.font.color.rgb = rgb(item_color)
        p.space_after = Pt(gap)
        p.line_spacing = 1.04
        if bullet:
            p.text = "•  " + p.text
    return box


def shape(slide, kind, x, y, w, h, fill=PANEL, line=None, radius=True):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1.1)
    return s


def card(slide, x, y, w, h, title, body, *, accent=CYAN, title_size=18, body_size=13):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, PANEL, GRID)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, 0.06, h, accent)
    add_text(slide, title, x + 0.22, y + 0.16, w - 0.38, 0.35, size=title_size, color=accent, bold=True)
    if isinstance(body, str):
        add_text(slide, body, x + 0.22, y + 0.62, w - 0.38, h - 0.76, size=body_size, color=MUTED)
    else:
        add_rich_lines(slide, body, x + 0.18, y + 0.59, w - 0.33, h - 0.7, size=body_size, color=MUTED, bullet=True)


def pill(slide, text, x, y, w, *, fill=PANEL_2, color=CYAN, size=11):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, 0.34, fill, None)
    add_text(slide, text, x, y + 0.01, w, 0.30, size=size, color=color, bold=True, align=PP_ALIGN.CENTER)


def connector(slide, x1, y1, x2, y2, *, color=CYAN, width=2.0, arrow=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if arrow:
        try:
            c.line.end_arrowhead = True
        except Exception:
            pass
    return c


def add_base(slide, title, kicker, page, source=None):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(BG)
    add_text(slide, kicker.upper(), 0.55, 0.30, 5.0, 0.25, size=9, color=CYAN, bold=True)
    add_text(slide, title, 0.55, 0.62, 12.1, 0.58, size=27, bold=True)
    shape(slide, MSO_SHAPE.RECTANGLE, 0.55, 1.25, 0.8, 0.035, CYAN)
    if source:
        add_text(slide, "依据：" + source, 0.56, 7.16, 11.8, 0.18, size=7.5, color="71829A")
    add_text(slide, f"{page:02d}", 12.25, 7.10, 0.5, 0.25, size=9, color="71829A", align=PP_ALIGN.RIGHT)


def section_slide(prs, number, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)
    shape(slide, MSO_SHAPE.OVAL, 0.7, 1.2, 1.25, 1.25, CYAN)
    add_text(slide, f"{number:02d}", 0.7, 1.46, 1.25, 0.55, size=30, color=BG, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, title, 2.25, 1.33, 9.9, 0.85, size=34, bold=True)
    add_text(slide, subtitle, 2.28, 2.28, 8.7, 0.8, size=17, color=MUTED)
    shape(slide, MSO_SHAPE.RECTANGLE, 2.28, 3.42, 8.8, 0.02, GRID)
    add_text(slide, "GENESIS WORLD · UAV · V2X", 2.28, 3.72, 6, 0.25, size=10, color=CYAN, bold=True)
    return slide


def add_table(slide, rows, cols, x, y, w, h, col_widths=None, font_size=12):
    table = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h)).table
    if col_widths:
        for i, width in enumerate(col_widths):
            table.columns[i].width = Inches(width)
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(PANEL_2 if r == 0 else PANEL)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            p = cell.text_frame.paragraphs[0]
            p.font.name = FONT
            p.font.size = Pt(font_size if r else font_size - 1)
            p.font.bold = r == 0
            p.font.color.rgb = rgb(CYAN if r == 0 else WHITE)
            p.vertical_anchor = MSO_ANCHOR.MIDDLE
    return table


def set_cell(table, row, col, text, *, color=None, bold=None, align=PP_ALIGN.LEFT):
    cell = table.cell(row, col)
    cell.text = text
    p = cell.text_frame.paragraphs[0]
    p.alignment = align
    p.font.name = FONT
    if bold is not None:
        p.font.bold = bold
    if color:
        p.font.color.rgb = rgb(color)


def build() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)

    # 1 — cover
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    if TEASER.exists():
        s.shapes.add_picture(str(TEASER), Inches(7.25), Inches(0), width=Inches(6.08), height=Inches(7.5))
        overlay = shape(s, MSO_SHAPE.RECTANGLE, 6.35, 0, 2.4, 7.5, BG)
        overlay.fill.transparency = 22
    if LOGO.exists():
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.55, 0.40, 3.55, 1.02, WHITE)
        s.shapes.add_picture(str(LOGO), Inches(0.67), Inches(0.55), width=Inches(3.25))
    pill(s, "课题组技术分享", 0.68, 1.55, 1.55, fill=PANEL_2, color=CYAN)
    add_text(s, "Genesis World", 0.67, 2.15, 6.35, 0.78, size=39, bold=True)
    add_text(s, "面向无人机与车联网研究的\n开放物理仿真平台", 0.67, 3.02, 6.2, 1.28, size=29, color=WHITE, bold=True)
    add_text(s, "从多物理与传感器，到四旋翼建模、强化学习与 V2X 联合仿真", 0.70, 4.58, 5.75, 0.75, size=16, color=MUTED)
    add_text(s, "源码基线  v1.3.3  ·  2026.09", 0.70, 6.60, 4.3, 0.30, size=11, color=CYAN, bold=True)

    # 2 — executive answer
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "先给结论：Genesis 应该放在研究栈的哪一层？", "核心判断", 2, "README.md；genesis/engine/scene.py")
    add_text(s, "物理世界 + 感知 + 高吞吐环境", 0.65, 1.56, 7.3, 0.55, size=29, color=CYAN, bold=True)
    add_text(s, "它是空—地智能体的“数字试验场”，不是一套完整的通信协议栈。", 0.66, 2.18, 8.1, 0.55, size=19, color=WHITE)
    card(s, 0.65, 3.06, 3.75, 2.55, "Genesis 原生强项", ["动力学、碰撞、场景", "RGB / Depth / LiDAR / IMU", "并行环境、控制与 RL 接口"], accent=GREEN, body_size=15)
    card(s, 4.78, 3.06, 3.75, 2.55, "需要外接", ["5G / NR-V2X / 802.11p", "包队列、调度、SINR / BER", "ns-3 / OMNeT++ / 信道模型"], accent=YELLOW, body_size=15)
    card(s, 8.91, 3.06, 3.75, 2.55, "课题组机会", ["空—地协同感知", "轨迹—通信联合优化", "多智能体 RL 与边缘卸载"], accent=PURPLE, body_size=15)

    # 3 — background
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "项目背景与定位", "01 · Genesis 概览", 3, "README.md；pyproject.toml；LICENSE")
    add_text(s, "Physical AI simulation platform", 0.65, 1.52, 6.1, 0.5, size=27, color=CYAN, bold=True)
    add_rich_lines(s, [
        "2024 年 12 月起源于学术项目，现由 Genesis AI 官方支持",
        "当前仓库包版本 1.3.3；Python 3.10–3.13；Apache 2.0",
        "统一多物理、真实感渲染、跨平台编译与 Python 接口",
        "目标：从单机原型扩展到数据中心 GPU 的 Physical AI 研发",
    ], 0.68, 2.18, 6.0, 2.7, size=16, color=MUTED, bullet=True, gap=10)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.2, 1.62, 5.45, 4.72, PANEL, GRID)
    for i, (num, label, detail, color) in enumerate([
        ("01", "统一", "同一场景、同一状态、多求解器", CYAN),
        ("02", "高吞吐", "批量环境与张量化控制", GREEN),
        ("03", "可扩展", "资产、传感器、控制器、算法", PURPLE),
        ("04", "跨平台", "CUDA / ROCm / Metal / CPU", YELLOW),
    ]):
        yy = 1.94 + i * 1.02
        pill(s, num, 7.55, yy, 0.55, fill=color, color=BG, size=10)
        add_text(s, label, 8.30, yy - 0.01, 1.3, 0.32, size=16, color=color, bold=True)
        add_text(s, detail, 9.42, yy, 2.75, 0.35, size=12.5, color=MUTED)

    # 4 — architecture
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "四层架构：把研究代码与硬件后端隔离", "01 · Genesis 概览", 4, "README.md — What is Genesis World?")
    layers = [
        ("研究应用", "控制 · 规划 · RL · 数据生成 · Agent", PURPLE),
        ("Simulation Interface", "Scene · Entity · Asset · Sensor · GUI", CYAN),
        ("Physics", "Rigid · FEM · MPM · PBD/SPH · uIPC · SAP", GREEN),
        ("Render", "Nyx · Luisa · Pyrender", BLUE),
        ("Compiler / Hardware", "Quadrants → CUDA · ROCm · Metal · Vulkan · x86 · ARM64", YELLOW),
    ]
    for i, (title, body, color) in enumerate(layers):
        yy = 1.52 + i * 0.99
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.04, yy, 11.25, 0.72, PANEL if i else PANEL_2, color)
        add_text(s, title, 1.34, yy + 0.17, 2.55, 0.28, size=15, color=color, bold=True)
        add_text(s, body, 4.0, yy + 0.17, 7.9, 0.28, size=14, color=WHITE)
        if i < len(layers) - 1:
            connector(s, 6.66, yy + 0.73, 6.66, yy + 0.97, color=GRID, width=1.5, arrow=False)

    # 5 — physics
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "统一多物理：不是只有刚体机器人", "01 · Genesis 概览", 5, "genesis/engine/simulator.py；examples/coupling/")
    items = [
        ("Rigid", "机器人、车辆、无人机、碰撞", CYAN),
        ("FEM", "弹性体与连续介质", GREEN),
        ("MPM", "沙、雪、材料大变形", YELLOW),
        ("PBD / SPH", "布料、液体、颗粒", BLUE),
        ("Stable Fluid", "烟雾/流体场", PURPLE),
        ("Couplers", "SAP · Legacy · uIPC", RED),
    ]
    for i, (title, body, color) in enumerate(items):
        x = 0.68 + (i % 3) * 4.18
        y = 1.58 + (i // 3) * 2.08
        card(s, x, y, 3.75, 1.68, title, body, accent=color, title_size=19, body_size=13)
    add_text(s, "研究含义：空中机器人、道路车辆、柔性物体、颗粒/流体可共享一个场景状态。", 0.72, 5.95, 11.9, 0.48, size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

    # 6 — scene lifecycle
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "一次仿真如何运行：Scene 是核心编排器", "02 · 实现细节", 6, "genesis/__init__.py；genesis/engine/scene.py；simulator.py")
    steps = [
        ("1", "gs.init", "后端 / 精度 / 随机种子"),
        ("2", "Scene", "求解器 / 渲染 / 传感器"),
        ("3", "add_entity", "Morph + Material + Surface"),
        ("4", "build", "解析 / 复制 / 分配 / 编译"),
        ("5", "step", "动作 → 物理 → 观测"),
    ]
    for i, (num, title, body) in enumerate(steps):
        x = 0.55 + i * 2.55
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.0, 2.05, 2.15, PANEL, CYAN if i in (0, 4) else GRID)
        shape(s, MSO_SHAPE.OVAL, x + 0.72, 1.63, 0.62, 0.62, CYAN if i < 4 else GREEN)
        add_text(s, num, x + 0.72, 1.78, 0.62, 0.24, size=14, color=BG, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, title, x + 0.16, 2.48, 1.73, 0.34, size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, body, x + 0.18, 3.02, 1.69, 0.64, size=12, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 4:
            connector(s, x + 2.06, 3.08, x + 2.50, 3.08, color=CYAN, width=2.0)
    add_text(s, "统一循环让传统控制、规划器和策略网络使用同一份状态与传感器数据。", 1.0, 5.13, 11.3, 0.55, size=19, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    # 7 — sensors
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "渲染与传感器：从“看起来像”到“可用于闭环”", "02 · 实现细节", 7, "README.md；genesis/engine/sensors/；options/sensors/options.py")
    card(s, 0.65, 1.55, 3.65, 2.1, "视觉", ["RGB / Camera", "Depth Camera", "Nyx · Luisa · Pyrender"], accent=BLUE, body_size=14)
    card(s, 4.52, 1.55, 3.65, 2.1, "机器人状态", ["IMU", "Joint Torque / Contact", "Surface Distance"], accent=GREEN, body_size=14)
    card(s, 8.39, 1.55, 3.65, 2.1, "环境感知", ["Raycaster / LiDAR", "Tactile", "Temperature"], accent=PURPLE, body_size=14)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.65, 4.08, 11.39, 1.62, PANEL_2, CYAN)
    add_text(s, "传感器时间模型", 0.98, 4.37, 2.25, 0.34, size=19, color=CYAN, bold=True)
    add_text(s, "history_length", 3.28, 4.29, 1.85, 0.33, size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, "历史观测", 3.28, 4.78, 1.85, 0.28, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(s, "delay", 5.56, 4.29, 1.42, 0.33, size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, "读取旧数据", 5.56, 4.78, 1.42, 0.28, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(s, "jitter", 7.38, 4.29, 1.42, 0.33, size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, "随机附加延迟", 7.38, 4.78, 1.42, 0.28, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(s, "≠ 网络协议栈", 9.32, 4.42, 2.1, 0.42, size=18, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)

    # 8 — parallel/RL
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "并行环境是 Genesis 对强化学习最直接的价值", "02 · 实现细节", 8, "parallel_simulation.py；domain_randomization.py；heterogeneous_simulation.py")
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.68, 1.62, 4.2, 4.6, PANEL, GRID)
    add_text(s, "scene.build(n_envs=B)", 0.98, 1.95, 3.62, 0.4, size=20, color=CYAN, bold=True, font=MONO, align=PP_ALIGN.CENTER)
    for row in range(3):
        for col in range(4):
            x = 1.04 + col * 0.86
            y = 2.72 + row * 0.75
            shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 0.62, 0.46, PANEL_2, GREEN if (row + col) % 3 == 0 else GRID)
            add_text(s, f"E{row*4+col+1}", x, y + 0.12, 0.62, 0.18, size=9, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, "张量化状态 / 动作 / reset", 1.05, 5.25, 3.45, 0.35, size=14, color=MUTED, align=PP_ALIGN.CENTER)
    card(s, 5.23, 1.62, 3.42, 2.02, "领域随机化", ["质量 / 质心 / 惯量", "摩擦 / 初值 / 目标", "传感器噪声与时间"], accent=GREEN, body_size=13)
    card(s, 8.92, 1.62, 3.42, 2.02, "异构环境", ["同批次不同几何", "不同环境独立参数", "只渲染少量环境"], accent=PURPLE, body_size=13)
    card(s, 5.23, 3.93, 7.11, 2.29, "边界", ["Genesis 提供环境与张量接口", "PPO 示例来自外部 rsl-rl", "SAC / MAPPO / QMIX 等算法需自行接入"], accent=YELLOW, body_size=14)

    # 9 — catalogue
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "仓库案例图谱：从物理验证到学习闭环", "03 · 具体案例", 9, "README.md；examples/")
    groups = [
        ("物理与耦合", "碰撞塔、布料、沙轮、液体、烟雾、刚柔耦合", GREEN),
        ("机器人控制", "Franka、Go2、IK、差分 IK、重力补偿", CYAN),
        ("感知", "RGB、深度、LiDAR、IMU、触觉、接触力", BLUE),
        ("学习", "无人机悬停 PPO、Go2、抓取、行为克隆", PURPLE),
        ("界面与工具", "GUI、键鼠交互、调试绘制、录制", YELLOW),
        ("本分支扩展", "移动机器人、开放词汇目标、RGB-D、A*", RED),
    ]
    for i, (title, body, color) in enumerate(groups):
        x = 0.68 + (i % 2) * 6.05
        y = 1.52 + (i // 2) * 1.57
        card(s, x, y, 5.68, 1.26, title, body, accent=color, title_size=16, body_size=12)
    add_text(s, "选择案例的原则：先验证物理与接口，再增加感知和算法复杂度。", 0.72, 6.40, 11.8, 0.35, size=16, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    section_slide(prs, 4, "无人机：从 URDF 到策略学习", "以仓库中的 Crazyflie 2.X 为主线，拆解模型、动力、控制与训练环境")

    # 11 — UAV examples
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "无人机案例全景", "04 · 无人机", 11, "examples/drone/")
    entries = [
        ("键盘飞行", "interactive_drone.py", "差分 RPM；理解执行器", CYAN),
        ("预定义轨迹", "fly.py --vis", "嵌入式 RPM 序列", BLUE),
        ("PID 航点", "fly_route.py", "位置→速度→姿态→mixer", GREEN),
        ("PPO 悬停", "hover_train/eval.py", "8192 环境；rsl-rl", PURPLE),
    ]
    for i, (title, cmd, body, color) in enumerate(entries):
        x = 0.70 + (i % 2) * 6.05
        y = 1.58 + (i // 2) * 2.23
        card(s, x, y, 5.65, 1.88, title, body, accent=color, title_size=20, body_size=13)
        add_text(s, cmd, x + 0.24, y + 1.30, 5.16, 0.30, size=11, color=WHITE, font=MONO)
    add_text(s, "由浅入深：执行器接口 → 稳定控制 → 点到点飞行 → 策略学习", 0.72, 6.20, 11.8, 0.42, size=18, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    # 12 — model
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "Crazyflie 2.X：模型由哪些部分构成？", "04 · 无人机", 12, "genesis/assets/urdf/drones/cf2x.urdf")
    # top view diagram
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.68, 1.52, 6.05, 4.98, PANEL, GRID)
    cx, cy = 3.70, 4.03
    # arms
    for x1, y1, x2, y2 in [(2.1, 2.55, 5.3, 5.5), (5.3, 2.55, 2.1, 5.5)]:
        connector(s, x1, y1, x2, y2, color="72829A", width=9, arrow=False)
    shape(s, MSO_SHAPE.OVAL, cx - 0.55, cy - 0.55, 1.1, 1.1, PANEL_2, CYAN)
    add_text(s, "27 g", cx - 0.55, cy - 0.10, 1.1, 0.25, size=14, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    rotors = [(2.05, 2.48, "P1", "+"), (5.15, 2.48, "P0", "−"), (2.05, 5.15, "P2", "−"), (5.15, 5.15, "P3", "+")]
    for x, y, name, spin in rotors:
        shape(s, MSO_SHAPE.OVAL, x, y, 0.58, 0.58, BG, GREEN if spin == "+" else PURPLE)
        add_text(s, name, x, y + 0.11, 0.58, 0.19, size=9, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, spin, x + 0.42, y - 0.08, 0.25, 0.20, size=12, color=GREEN if spin == "+" else PURPLE, bold=True)
    add_text(s, "旋翼力臂： (±0.028, ±0.028, 0) m", 1.30, 5.90, 4.85, 0.28, size=12, color=MUTED, align=PP_ALIGN.CENTER)
    # facts
    facts = [
        ("质量", "0.027 kg"),
        ("惯量", "Ixx=Iyy=1.4e−5\nIzz=2.17e−5 kg·m²"),
        ("推力系数", "KF = 3.16e−10"),
        ("反扭矩系数", "KM = 7.94e−12"),
        ("碰撞近似", "Ø 0.12 m × 0.025 m 圆柱"),
    ]
    for i, (label, value) in enumerate(facts):
        yy = 1.58 + i * 0.95
        add_text(s, label, 7.15, yy + 0.12, 1.25, 0.28, size=13, color=MUTED)
        add_text(s, value, 8.53, yy + 0.02, 3.95, 0.52, size=16 if i != 1 else 13, color=WHITE, bold=True)

    # 13 — rotor dynamics
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "旋翼动力如何进入刚体求解器？", "04 · 无人机", 13, "drone_entity.py；rigid/abd/accessor.py")
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.70, 1.60, 5.35, 3.02, PANEL_2, CYAN)
    add_text(s, "Fᵢ = [0, 0, KF · RPMᵢ²]ᵀ", 1.05, 2.12, 4.65, 0.52, size=24, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, "Mᵢ = [0, 0, sᵢ KM · RPMᵢ²]ᵀ", 1.05, 3.03, 4.65, 0.52, size=23, color=PURPLE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, "局部坐标施力 + 旋翼位置力臂 → roll / pitch / yaw", 0.98, 4.07, 4.82, 0.30, size=12.5, color=MUTED, align=PP_ALIGN.CENTER)
    card(s, 6.42, 1.60, 5.92, 1.33, "RPM 接口", "每步调用一次；支持 (4,) 与 (B,4)", accent=GREEN, title_size=17, body_size=14)
    card(s, 6.42, 3.17, 5.92, 1.45, "悬停转速", "√(mg / 4KF) ≈ 14,469 RPM", accent=YELLOW, title_size=17, body_size=18)
    add_text(s, "这是低阶旋翼模型，不是 CFD；风场、下洗和地面效应需要扩展。", 0.73, 5.38, 11.55, 0.58, size=18, color=RED, bold=True, align=PP_ALIGN.CENTER)

    # 14 — PID
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "PID 航点案例：一个可解释的飞行控制基线", "04 · 无人机", 14, "examples/drone/fly_route.py；quadcopter_controller.py")
    flow = [
        ("目标点", "x*, y*, z*", PURPLE),
        ("位置环", "误差 → 期望速度", CYAN),
        ("速度环", "速度误差 → 推力/倾斜", GREEN),
        ("姿态环", "roll / pitch / yaw", BLUE),
        ("Mixer", "M1…M4 RPM", YELLOW),
    ]
    for i, (title, body, color) in enumerate(flow):
        x = 0.55 + i * 2.55
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.03, 2.05, 1.48, PANEL, color)
        add_text(s, title, x + 0.13, 2.33, 1.79, 0.3, size=17, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, body, x + 0.13, 2.88, 1.79, 0.31, size=11.5, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 4:
            connector(s, x + 2.06, 2.77, x + 2.48, 2.77, color=CYAN, width=2.0)
    card(s, 0.72, 4.35, 3.70, 1.50, "航点", "(1,1,2) → (−1,2,1) → (0,0,0.5)", accent=PURPLE, body_size=13)
    card(s, 4.82, 4.35, 3.70, 1.50, "停止条件", "距离 < 0.1 m，或单点 1000 步", accent=GREEN, body_size=13)
    card(s, 8.92, 4.35, 3.70, 1.50, "输出", "out/fly_route.mp4", accent=YELLOW, body_size=13)
    add_text(s, "局限：点到点控制 ≠ 三维避障与动力学可行轨迹规划", 0.72, 6.28, 11.8, 0.36, size=16, color=RED, bold=True, align=PP_ALIGN.CENTER)

    # 15 — PPO
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "PPO 悬停：环境、策略与奖励如何闭环", "04 · 无人机", 15, "examples/drone/hover_env.py；hover_train.py")
    # loop
    cards = [
        ("8192× 环境", "100 Hz", CYAN, 0.70),
        ("17 维观测", "相对位置 / 姿态 / 速度 / 上一动作", BLUE, 3.83),
        ("PPO Actor", "128×128 tanh", PURPLE, 7.00),
        ("4 维动作", "RPM 扰动", GREEN, 10.18),
    ]
    for title, body, color, x in cards:
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.78, 2.46, 1.37, PANEL, color)
        add_text(s, title, x + 0.12, 2.02, 2.22, 0.28, size=16, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, body, x + 0.12, 2.50, 2.22, 0.35, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
    for x in (3.18, 6.35, 9.52):
        connector(s, x, 2.47, x + 0.56, 2.47, color=CYAN, width=2)
    connector(s, 11.40, 3.18, 2.00, 3.70, color=GREEN, width=1.5)
    add_text(s, "reward / done / next obs", 5.08, 3.42, 3.3, 0.28, size=10.5, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    card(s, 0.70, 4.13, 4.00, 1.72, "奖励", ["接近目标 +", "动作抖动 / 角速度 −", "坠毁 −10"], accent=GREEN, body_size=12.5)
    card(s, 4.95, 4.13, 3.52, 1.72, "外部依赖", ["rsl-rl-lib ≥ 5", "PPO 不是 Genesis 内置", "TensorBoard 日志"], accent=YELLOW, body_size=12.5)
    card(s, 8.72, 4.13, 3.91, 1.72, "实现提醒", ["同名 logs 会被删除", "action latency 开关未生效", "先小批量再 8192"], accent=RED, body_size=12.5)

    # 16 — path planning
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "从“能飞”到“会规划”：推荐分层，而非端到端堆叠", "05 · 规划与扩展", 16, "fly_route.py；规划模块为建议架构")
    layers = [
        ("任务层", "语义目标 · 任务分配 · 航点序列", PURPLE),
        ("全局规划", "3D A* / RRT* / Kinodynamic planning", CYAN),
        ("局部安全", "Depth / LiDAR · 局部地图 · 动态重规划", GREEN),
        ("轨迹控制", "MPC / PID · 速度与倾角约束", BLUE),
        ("执行器", "Mixer → set_propellers_rpm()", YELLOW),
    ]
    for i, (title, body, color) in enumerate(layers):
        yy = 1.53 + i * 0.95
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.15, yy, 11.0, 0.65, PANEL, color)
        add_text(s, title, 1.47, yy + 0.16, 1.65, 0.25, size=14, color=color, bold=True)
        add_text(s, body, 3.30, yy + 0.16, 8.45, 0.25, size=13, color=WHITE)
    add_text(s, "学习算法优先替换局部策略或控制残差；全局规划和安全过滤保留可解释基线。", 0.83, 6.34, 11.65, 0.35, size=16, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    # 17 — local mobile robot
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "本仓库扩展案例：自然语言目标 → 几何安全路径", "05 · 规划与扩展", 17, "examples/mobile_robot/room_navigation_vision.py（非上游官方案例）")
    if ROBOT_IMAGE.exists():
        s.shapes.add_picture(str(ROBOT_IMAGE), Inches(0.70), Inches(1.60), width=Inches(5.05), height=Inches(3.79))
        shape(s, MSO_SHAPE.RECTANGLE, 0.70, 5.39, 5.05, 0.50, PANEL_2)
        add_text(s, "本地 RGB 视角：语义目标与障碍场景", 0.84, 5.52, 4.77, 0.23, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    stages = [
        ("文本", "开放词汇检测", PURPLE),
        ("RGB-D", "目标世界坐标", BLUE),
        ("A*", "障碍膨胀 + 禁止切角", GREEN),
        ("安全层", "航点 + LiDAR", YELLOW),
    ]
    for i, (title, body, color) in enumerate(stages):
        yy = 1.62 + i * 1.05
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 6.25, yy, 5.85, 0.78, PANEL, color)
        add_text(s, title, 6.50, yy + 0.20, 1.15, 0.26, size=15, color=color, bold=True)
        add_text(s, body, 7.83, yy + 0.20, 3.95, 0.26, size=13, color=WHITE)
    add_text(s, "可迁移：2D 栅格 → 3D 体素；车体膨胀 → UAV 安全球/椭球", 6.35, 6.15, 5.60, 0.44, size=15, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    section_slide(prs, 6, "V2X：明确边界，再做联合仿真", "Genesis 负责物理和感知，网络模拟器负责协议、链路、队列与调度")

    # 19 — capability matrix
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "通信能力矩阵：哪些能做，哪些不能直接宣称？", "06 · V2X 联合仿真", 19, "仓库全文检索；SensorOptions delay/jitter")
    data = [
        ("能力", "现状", "处理方式"),
        ("位姿 / 动力学 / 遮挡", "原生", "Genesis"),
        ("RGB / Depth / LiDAR / IMU", "原生", "Genesis Sensor API"),
        ("观测 delay / jitter / history", "部分原生", "基础时延实验"),
        ("包队列 / 拥塞 / 吞吐", "未发现", "ns-3 / OMNeT++"),
        ("V2X 协议栈 / 调度", "未发现", "专业网络仿真器"),
        ("SINR / BER / MCS", "未发现完整模型", "信道与链路模型"),
    ]
    table = add_table(s, len(data), 3, 0.72, 1.58, 11.90, 4.95, col_widths=[4.4, 2.3, 5.2], font_size=12.5)
    for r, row in enumerate(data):
        for c, value in enumerate(row):
            set_cell(table, r, c, value, color=CYAN if r == 0 else WHITE, bold=r == 0)
        if r > 0:
            status_color = GREEN if row[1] == "原生" else YELLOW if row[1] == "部分原生" else RED
            set_cell(table, r, 1, row[1], color=status_color, bold=True, align=PP_ALIGN.CENTER)

    # 20 — co-sim architecture
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "推荐架构：物理—网络—决策三层锁步闭环", "06 · V2X 联合仿真", 20, "建议方案；并非 Genesis 内置 V2X")
    blocks = [
        ("Genesis World", "UAV / Vehicle\nPhysics · Scene · Sensor", CYAN, 0.70),
        ("同步桥", "sim_time · ID · pose\nqueue · timeout · seed", YELLOW, 4.62),
        ("Network Simulator", "ns-3 / OMNeT++\nlatency · loss · SINR", PURPLE, 8.54),
    ]
    for title, body, color, x in blocks:
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.75, 3.35, 2.07, PANEL, color)
        add_text(s, title, x + 0.18, 2.05, 2.99, 0.35, size=18, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, body, x + 0.20, 2.62, 2.95, 0.68, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    connector(s, 4.07, 2.77, 4.55, 2.77, color=GREEN, width=2)
    connector(s, 8.00, 2.77, 8.47, 2.77, color=GREEN, width=2)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 2.45, 4.65, 8.43, 1.30, PANEL_2, GREEN)
    add_text(s, "规划 / 控制 / MARL / 边缘卸载", 2.78, 4.94, 7.78, 0.35, size=21, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    connector(s, 6.66, 3.85, 6.66, 4.58, color=GREEN, width=2)
    add_text(s, "dt_phys = 0.01 s；dt_net = N × dt_phys；消息携带时间戳并按到达时刻释放", 1.15, 6.32, 11.0, 0.36, size=14.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

    # 21 — research topics
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "三类可形成论文闭环的课题", "07 · 研究落地", 21, "建议研究设计")
    card(s, 0.68, 1.60, 3.75, 4.40, "01  空—地协同感知", [
        "UAV 高空观测事故/目标",
        "车辆融合本地 LiDAR 与空中消息",
        "变量：定位误差、时延、丢包、遮挡",
        "指标：成功率、碰撞率、AoI",
    ], accent=CYAN, title_size=19, body_size=13.5)
    card(s, 4.79, 1.60, 3.75, 4.40, "02  移动中继", [
        "多 UAV 作为 V2X 中继",
        "联合优化航点、功率、信道与关联",
        "MARL：MAPPO / MADDPG / QMIX",
        "指标：覆盖、吞吐、能耗、公平性",
    ], accent=PURPLE, title_size=19, body_size=13.5)
    card(s, 8.90, 1.60, 3.75, 4.40, "03  边缘推理卸载", [
        "传关键帧、目标框或中间特征",
        "比较本地、边缘与协同推理",
        "通信约束进入状态与奖励",
        "指标：精度、时延、带宽、能耗",
    ], accent=GREEN, title_size=19, body_size=13.5)
    add_text(s, "共同原则：网络结果必须改变观测、动作或奖励，不能只作为离线日志。", 0.83, 6.35, 11.65, 0.35, size=16, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)

    # 22 — roadmap
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "建议落地路线：12 周形成可复现实验链", "07 · 研究落地", 22, "建议实施计划")
    phases = [
        ("0", "1–2 周", "基线复现", "跑通 UAV 案例\n版本与碰撞验证", CYAN),
        ("1", "2–4 周", "三维导航", "Depth/LiDAR/IMU\n3D A* / RRT*", GREEN),
        ("2", "3–6 周", "空—地多智能体", "统一 ID / 时钟 / 日志\nPython channel model", BLUE),
        ("3", "6–12 周", "网络 + MARL", "接入 ns-3/OMNeT++\n联合优化与消融", PURPLE),
    ]
    for i, (num, duration, title, body, color) in enumerate(phases):
        x = 0.65 + i * 3.12
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.05, 2.65, 3.42, PANEL, color)
        shape(s, MSO_SHAPE.OVAL, x + 0.87, 1.61, 0.90, 0.90, color)
        add_text(s, num, x + 0.87, 1.84, 0.90, 0.32, size=18, color=BG, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, duration, x + 0.28, 2.72, 2.09, 0.27, size=11, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, title, x + 0.20, 3.20, 2.25, 0.35, size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, body, x + 0.22, 3.88, 2.20, 0.82, size=12.5, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 3:
            connector(s, x + 2.67, 3.75, x + 3.06, 3.75, color=CYAN, width=2)
    add_text(s, "每阶段都保留可运行基线、固定随机种子、统一指标与可回放日志。", 0.76, 6.24, 11.80, 0.40, size=17, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    # 23 — risks
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "采用前必须正视的限制", "07 · 研究落地", 23, "当前源码审阅结论")
    risks = [
        ("API 演进", "锁版本与提交，不只写 latest", YELLOW),
        ("气动简化", "RPM² 集中力；非 CFD", RED),
        ("碰撞待验证", "源码说明与 URDF collision 并存", RED),
        ("训练脚本", "同名日志覆盖；动作时延开关未生效", YELLOW),
        ("平台差异", "GPU / 渲染 / 批量传感器分别测", BLUE),
        ("通信缺口", "协议与链路必须外接", PURPLE),
    ]
    for i, (title, body, color) in enumerate(risks):
        x = 0.68 + (i % 2) * 6.03
        y = 1.55 + (i // 2) * 1.53
        card(s, x, y, 5.63, 1.20, title, body, accent=color, title_size=15, body_size=11.5)
    add_text(s, "工程判断：Genesis 值得作为研究底座，但实验结论必须经过任务级校准和验证。", 0.83, 6.33, 11.65, 0.35, size=16, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    # 24 — benchmark
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "建议第一个组内基准任务", "08 · 行动建议", 24, "文档附录 B")
    add_text(s, "通信受限的空—地协同目标搜索与安全到达", 0.72, 1.56, 11.8, 0.52, size=27, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    blocks = [
        ("场景", "1 UAV + 2–8 车辆\n随机目标与障碍", BLUE),
        ("通信档位", "理想 / 固定损伤\n/ ns-3", PURPLE),
        ("算法基线", "A* + 安全层\n/ MAPPO", GREEN),
        ("评价", "成功 / 碰撞 / AoI\n吞吐 / 能耗 / RTF", YELLOW),
    ]
    for i, (title, body, color) in enumerate(blocks):
        x = 0.68 + i * 3.13
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.55, 2.68, 2.37, PANEL, color)
        add_text(s, title, x + 0.18, 2.92, 2.32, 0.32, size=17, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, body, x + 0.18, 3.55, 2.32, 0.72, size=13, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 3:
            connector(s, x + 2.71, 3.70, x + 3.05, 3.70, color=CYAN, width=2)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.35, 5.62, 10.63, 0.74, PANEL_2, CYAN)
    add_text(s, "一次覆盖无人机 · 车联网 · 路径规划 · 强化学习，并可清楚做模块消融", 1.55, 5.85, 10.23, 0.28, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

    # 25 — conclusion
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    if LOGO.exists():
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.55, 0.41, 3.55, 1.00, WHITE)
        s.shapes.add_picture(str(LOGO), Inches(0.70), Inches(0.58), width=Inches(3.2))
    add_text(s, "结论", 0.70, 1.78, 2.2, 0.55, size=34, color=CYAN, bold=True)
    add_text(s, "Genesis 的价值，不是把所有工具合并成一个，\n而是把物理、感知与学习闭环做成可扩展的共同底座。", 0.72, 2.65, 11.6, 1.35, size=27, color=WHITE, bold=True)
    add_text(s, "Genesis：物理与传感器  ×  网络仿真器：通信真实性  ×  算法框架：智能决策", 0.74, 4.52, 11.55, 0.48, size=18, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    shape(s, MSO_SHAPE.RECTANGLE, 0.72, 5.55, 11.88, 0.025, GRID)
    add_text(s, "下一步：复现 UAV 基线 → 验证碰撞 → 建立三维导航 → 接入网络联合仿真", 0.78, 5.94, 11.43, 0.48, size=17, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(s, "Q & A", 10.85, 6.78, 1.45, 0.35, size=18, color=CYAN, bold=True, align=PP_ALIGN.RIGHT)

    return prs


if __name__ == "__main__":
    presentation = build()
    presentation.save(OUT)
    print(f"Wrote {OUT} ({len(presentation.slides)} slides)")
